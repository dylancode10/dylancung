"""Load LinkedIn posts from the official data export.

Two inputs are supported:

* ``Shares.csv`` from *Settings -> Data privacy -> Get a copy of your data*
  (the text, date and link of every post you wrote). Pass the CSV itself, the
  folder you unzipped the export into, or the export ``.zip``.
* Optionally, the Creator analytics ``.xlsx`` from *Analytics -> Content ->
  Export*. Its ``TOP POSTS`` sheet adds impressions and engagements per post.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

HASHTAG_RE = re.compile(r"#(\w+)")
URL_RE = re.compile(r"https?://\S+")
POST_ID_RE = re.compile(r"(\d{15,})")


@dataclass
class Post:
    id: str
    date: datetime
    text: str
    url: str = ""
    shared_url: str = ""
    media_url: str = ""
    impressions: int | None = None
    engagements: int | None = None
    hashtags: list[str] = field(default_factory=list)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    @property
    def has_media(self) -> bool:
        return bool(self.media_url)

    @property
    def has_link(self) -> bool:
        return bool(self.shared_url) or bool(URL_RE.search(self.text))

    @property
    def engagement_rate(self) -> float | None:
        if self.impressions and self.engagements is not None:
            return self.engagements / self.impressions
        return None


def _post_id(url: str) -> str:
    match = POST_ID_RE.search(unquote(url or ""))
    return match.group(1) if match else ""


def _parse_date(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    value = str(value or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def _read_shares_csv_text(path: Path) -> str:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            name = next((n for n in zf.namelist() if n.lower().endswith("shares.csv")), None)
            if name is None:
                raise FileNotFoundError(f"No Shares.csv inside {path}")
            return zf.read(name).decode("utf-8-sig")
    if path.is_dir():
        matches = sorted(path.rglob("Shares.csv"))
        if not matches:
            raise FileNotFoundError(f"No Shares.csv under {path}")
        path = matches[0]
    return path.read_text(encoding="utf-8-sig")


def load_shares(path: str | Path) -> list[Post]:
    """Parse posts from Shares.csv (or an export folder / zip containing it)."""
    reader = csv.DictReader(io.StringIO(_read_shares_csv_text(Path(path))))
    posts: list[Post] = []
    for i, row in enumerate(reader):
        date = _parse_date(row.get("Date"))
        if date is None:
            continue
        text = (row.get("ShareCommentary") or "").strip().strip('"').replace('""', '"')
        url = (row.get("ShareLink") or "").strip()
        posts.append(
            Post(
                id=_post_id(url) or f"row{i + 1}",
                date=date,
                text=text,
                url=url,
                shared_url=(row.get("SharedUrl") or "").strip(),
                media_url=(row.get("MediaUrl") or "").strip(),
                hashtags=[h.lower() for h in HASHTAG_RE.findall(text)],
            )
        )
    posts.sort(key=lambda p: p.date)
    return posts


def load_analytics(path: str | Path) -> list[dict]:
    """Read per-post metrics from the TOP POSTS sheet of a Creator analytics export.

    The sheet holds two side-by-side tables (top posts by engagements and by
    impressions), each with columns ``Post URL | Post publish date | <metric>``.
    Returns one dict per post URL: ``{"id", "url", "date", "impressions", "engagements"}``.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = next((wb[n] for n in wb.sheetnames if n.strip().upper() == "TOP POSTS"), None)
    if sheet is None:
        raise ValueError(f"No 'TOP POSTS' sheet in {path}; sheets: {wb.sheetnames}")

    tables: list[tuple[int, str]] = []  # (url column index, metric name)
    metrics: dict[str, dict] = {}
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if c is None else c for c in row]
        if not tables:
            for i, cell in enumerate(cells):
                if str(cell).strip().lower() == "post url" and i + 2 < len(cells):
                    tables.append((i, str(cells[i + 2]).strip().lower()))
            continue
        for col, metric in tables:
            url = str(cells[col]).strip() if col < len(cells) else ""
            if not url.startswith("http"):
                continue
            entry = metrics.setdefault(
                url, {"id": _post_id(url), "url": url, "date": _parse_date(cells[col + 1])}
            )
            try:
                value = int(float(cells[col + 2]))
            except (TypeError, ValueError):
                continue
            if "impression" in metric:
                entry["impressions"] = value
            elif "engagement" in metric:
                entry["engagements"] = value
    return list(metrics.values())


def merge_analytics(posts: list[Post], analytics: list[dict]) -> int:
    """Attach metrics to posts. Matches on post id, then on a unique same-day post.

    Returns the number of posts that received metrics.
    """
    by_id = {p.id: p for p in posts}
    by_day: dict = {}
    for p in posts:
        by_day.setdefault(p.date.date(), []).append(p)

    matched = 0
    for entry in analytics:
        post = by_id.get(entry["id"])
        if post is None and entry.get("date"):
            same_day = by_day.get(entry["date"].date(), [])
            post = same_day[0] if len(same_day) == 1 else None
        if post is None:
            continue
        post.impressions = entry.get("impressions", post.impressions)
        post.engagements = entry.get("engagements", post.engagements)
        matched += 1
    return matched
